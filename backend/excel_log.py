"""
Excel waiting list manager for routine patient cases.
Uses openpyxl to create/update waiting-list.xlsx.
"""

import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_FILE = os.path.join(os.path.dirname(__file__), "waiting-list.xlsx")

HEADERS = [
    "Ticket No.", "Date", "Time", "Patient Name", "Age", "Gender",
    "Phone", "Email", "Chief Complaint", "Duration", "Pain Scale",
    "Consciousness", "Medical History", "Medications", "Allergies",
    "Triage Level", "AI Reasoning", "Recommendations"
]

HEADER_FILL = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
ROUTINE_FILL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
URGENT_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)


def _create_workbook():
    """Create a new workbook with formatted headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Waiting List"
    ws.freeze_panes = "A2"

    # Column widths
    col_widths = [16, 12, 10, 22, 6, 10, 14, 28, 35, 16, 12, 14, 28, 28, 22, 14, 50, 50]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Write headers
    ws.row_dimensions[1].height = 30
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    return wb, ws


def _get_or_create_workbook():
    """Open existing workbook or create a new one."""
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
    else:
        wb, ws = _create_workbook()
    return wb, ws


def _get_next_ticket_number(ws) -> str:
    """Generate a ticket number based on today's date and row count."""
    today = datetime.now().strftime("%Y%m%d")
    max_row = ws.max_row
    # Count rows for today
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[0]).startswith(f"RHC-{today}"):
            count += 1
    return f"RHC-{today}-{count + 1:04d}"


def add_to_waiting_list(patient_data: dict, triage_result: dict) -> dict:
    """
    Add a patient to the waiting list Excel file.
    Returns {'success': bool, 'ticket_number': str, 'error': str}
    """
    try:
        wb, ws = _get_or_create_workbook()
        ticket = _get_next_ticket_number(ws)
        now = datetime.now()
        level = triage_result.get("level", "ROUTINE")

        row_data = [
            ticket,
            now.strftime("%d-%m-%Y"),
            now.strftime("%I:%M %p"),
            patient_data.get("name", ""),
            patient_data.get("age", ""),
            patient_data.get("gender", ""),
            patient_data.get("phone", ""),
            patient_data.get("email", ""),
            patient_data.get("symptoms", ""),
            patient_data.get("duration", ""),
            patient_data.get("pain_scale", ""),
            patient_data.get("consciousness", ""),
            patient_data.get("medical_history", ""),
            patient_data.get("medications", ""),
            patient_data.get("allergies", ""),
            level,
            triage_result.get("reasoning", ""),
            triage_result.get("recommendations", ""),
        ]

        next_row = ws.max_row + 1
        fill = ROUTINE_FILL if level == "ROUTINE" else URGENT_FILL if level == "URGENT" else ALT_FILL
        if next_row % 2 == 0 and level == "ROUTINE":
            fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=next_row, column=col, value=value)
            cell.border = THIN_BORDER
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col == 1:
                cell.font = Font(bold=True, color="1E40AF")
                cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[next_row].height = 25
        wb.save(EXCEL_FILE)
        return {"success": True, "ticket_number": ticket, "error": None}

    except Exception as e:
        return {"success": False, "ticket_number": None, "error": str(e)}


def get_waiting_list() -> list:
    """Return all rows from the waiting list as a list of dicts."""
    if not os.path.exists(EXCEL_FILE):
        return []
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(cell is not None for cell in row):
                rows.append(dict(zip(HEADERS, row)))
        return rows
    except Exception:
        return []
